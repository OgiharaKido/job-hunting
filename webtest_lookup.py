#!/usr/bin/env python3
import argparse
import json
import re
import unicodedata
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path("/Users/kido/Desktop/就活")
XLSX_PATH = ROOT / "webテスト(神） のコピー.xlsx"
CACHE_PATH = ROOT / ".webtest_index.jsonl"

HEADER_HINTS = {"解答", "回答", "解", "答え", "設問", "問題", "問題文", "現在", "確認", "メモ"}


def normalize(text: str) -> str:
    text = unicodedata.normalize("NFKC", text).lower()
    text = re.sub(r"\s+", " ", text)
    return text.strip()


def stringify(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()


def build_index() -> int:
    wb = load_workbook(XLSX_PATH, read_only=True, data_only=True)
    count = 0
    with CACHE_PATH.open("w", encoding="utf-8") as fh:
        for ws in wb.worksheets:
            header = [stringify(v) for v in next(ws.iter_rows(values_only=True), ())]
            for row_no, row in enumerate(ws.iter_rows(values_only=True), start=1):
                cells = [stringify(v) for v in row]
                if not any(cells):
                    continue
                if row_no == 1 and sum(1 for c in cells if c in HEADER_HINTS) >= 1:
                    continue
                nonempty = [(idx, cell) for idx, cell in enumerate(cells) if cell]
                if not nonempty:
                    continue
                answer = ""
                for idx, name in enumerate(header):
                    if name in {"解答", "回答", "解", "答え"} and idx < len(cells) and cells[idx]:
                        answer = cells[idx]
                        break
                if not answer and len(nonempty) >= 2:
                    answer = nonempty[1][1]
                searchable = " ".join(cells)
                record = {
                    "sheet": ws.title,
                    "row": row_no,
                    "answer": answer,
                    "cells": [cell for _, cell in nonempty[:6]],
                    "search": normalize(searchable),
                }
                fh.write(json.dumps(record, ensure_ascii=False) + "\n")
                count += 1
    return count


def search_index(query: str, limit: int) -> list[dict]:
    tokens = [t for t in normalize(query).split(" ") if t]
    results = []
    with CACHE_PATH.open(encoding="utf-8") as fh:
        for line in fh:
            rec = json.loads(line)
            score = 0
            for token in tokens:
                if token in rec["search"]:
                    score += len(token) * 2
            if not score:
                continue
            score += max(0, 10 - abs(len(rec["search"]) - len(normalize(query))))
            rec["score"] = score
            results.append(rec)
    results.sort(key=lambda r: (-r["score"], r["sheet"], r["row"]))
    return results[:limit]


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("query", nargs="?", help="keyword(s) to search")
    parser.add_argument("--rebuild", action="store_true")
    parser.add_argument("--limit", type=int, default=5)
    args = parser.parse_args()

    if args.rebuild or not CACHE_PATH.exists():
        count = build_index()
        if not args.query:
            print(f"indexed {count} rows into {CACHE_PATH}")
            return

    if not args.query:
        parser.error("query is required unless --rebuild is used alone")

    for rec in search_index(args.query, args.limit):
        cells = " | ".join(rec["cells"])
        print(f"[{rec['sheet']} row {rec['row']}] answer={rec['answer']}")
        print(cells)
        print()


if __name__ == "__main__":
    main()
